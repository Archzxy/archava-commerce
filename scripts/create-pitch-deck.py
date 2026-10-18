#!/usr/bin/env python3
"""Populate a downloaded organizer template; never edit the online master.

Requires python-pptx and the reviewed, cropped public demo images described in
docs/PITCH_DECK.md. LibreOffice exports the generated editable PPTX separately.
"""

from argparse import ArgumentParser
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PURPLE = '7B3DFF'
DARK = '1F1F24'
MUTED = '625A6B'
CREAM = 'FFF9F5'
LILAC = 'F1E9FF'
LIME = 'D5FF4F'
PINK = 'FFE6F4'
WHITE = 'FFFFFF'
LINE = 'E6DCEE'
FONT = 'Inter'
URL = 'https://dopamine-656573292283.asia-southeast1.run.app'
REPO = 'https://github.com/Archzxy/archava-commerce'
VIDEO = 'https://drive.google.com/drive/folders/1ZFEBzCy2zcdvwWMdTqUovlTCvuU9ODZp?usp=sharing'


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=CREAM, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, x, y, w, h, value, size=14, color=DARK, bold=False,
         align=PP_ALIGN.LEFT, link=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.12
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if link:
            run.hyperlink.address = link
    return shape


def arrow(slide, x1, y1, x2, y2, both=False, color=PURPLE):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.6)
    ln = c._element.spPr.get_or_add_ln()
    end = OxmlElement('a:tailEnd')
    end.set('type', 'triangle')
    ln.append(end)
    if both:
        begin = OxmlElement('a:headEnd')
        begin.set('type', 'triangle')
        ln.append(begin)
    return c


def picture(slide, path, x, y, max_w, max_h):
    from PIL import Image
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (max_w-w)/2), Inches(y + (max_h-h)/2), Inches(w), Inches(h))


def pill(slide, x, y, w, label, fill=LIME, color=DARK):
    rect(slide, x, y, w, 0.31, fill, rounded=True)
    text(slide, x + 0.09, y + 0.065, w - 0.18, 0.2, label, 9.5, color, True)


def clean(slide):
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            shape._element.getparent().remove(shape._element)


def setup(slide, title, index, assets, subtitle=None):
    clean(slide)
    text(slide, 0.34, 0.76, 9.25, 0.46, title, 20, DARK, True)
    rect(slide, 0.35, 1.30, 0.54, 0.045, PURPLE)
    if subtitle:
        text(slide, 1.03, 1.24, 8.45, 0.24, subtitle, 10, MUTED)
    text(slide, 0.35, 5.29, 8.3, 0.16, 'DOPAMINE  /  ARCHZXY  /  PROTOTYPE  /  18 OCTOBER 2026', 7.8, MUTED)
    text(slide, 9.1, 5.25, 0.5, 0.2, f'{index:02}', 10, PURPLE, True, PP_ALIGN.RIGHT)


def note(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def remove_slide(prs, slide):
    for sid in list(prs.slides._sldIdLst):
        if prs.part.related_slide(sid.rId) is slide:
            prs.part.drop_rel(sid.rId)
            prs.slides._sldIdLst.remove(sid)
            return
    raise ValueError('Slide not found')


def main():
    parser = ArgumentParser(description=__doc__)
    parser.add_argument('--template', type=Path, default=ROOT / '.logs/pitch-pdf-20261018/official-template.pptx')
    parser.add_argument('--assets', type=Path, default=ROOT / '.logs/pitch-pdf-20261018/assets')
    parser.add_argument('--output', type=Path, default=ROOT / 'output/submission/dopamine-2026-10-19')
    args = parser.parse_args()
    assets = args.assets
    prs = Presentation(args.template)
    assert len(prs.slides) == 16, 'Expected the reviewed sixteen-page organizer template'
    original = list(prs.slides)

    # A second screenshot page repeats the exact organizer section/layout.
    # Allocate its part/notes before removal: python-pptx numbers a new slide
    # from collection length, which otherwise collides with retained parts.
    screenshot_extra = prs.slides.add_slide(original[10].slide_layout)
    screenshot_extra.notes_slide
    for sh in list(screenshot_extra.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in original[10].shapes:
        if sh.has_text_frame:
            continue
        copy = deepcopy(sh._element)
        if sh.shape_type == 13:  # Rebind the copied organizer header/footer media.
            embed = copy.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
            if embed is not None:
                rel_id = embed.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                relation = original[10].part.rels[rel_id]
                new_id = screenshot_extra.part.relate_to(relation.target_part, relation.reltype)
                embed.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', new_id)
        screenshot_extra.shapes._spTree.insert_element_before(copy, 'p:extLst')
    # Remove only the download instructions, two explicitly optional sections,
    # and the unused blank slide. All required sections and the closing remain.
    for i in [1, 6, 9, 14]:
        remove_slide(prs, original[i])
    slide_ids = prs.slides._sldIdLst
    new_id = slide_ids[-1]
    slide_ids.remove(new_id)
    screenshot_pos = next(i for i, sl in enumerate(prs.slides) if sl is original[10])
    slide_ids.insert(screenshot_pos + 1, new_id)
    numbering = {id(sl): i+1 for i, sl in enumerate(prs.slides)}

    # Team details retain the original full-size organizer banner.
    s = original[0]
    clean(s)
    text(s, 0.35, 3.10, 6.3, 0.35, 'Team Details', 18, DARK, True)
    text(s, 0.36, 3.61, 6.4, 0.23, 'TEAM  Archzxy', 12, PURPLE, True)
    text(s, 0.36, 3.99, 6.5, 0.28, 'Team leader: Muhammad Fikri Haikal', 12.5)
    text(s, 0.36, 4.32, 6.5, 0.25, 'Team member: Muhammad Afif Alfawwaz', 12.5)
    text(s, 0.36, 4.75, 6.6, 0.53, 'Problem Statement: Retail & Commerce:\nIntelligent Customer and Business Experiences', 11.5, MUTED)
    picture(s, assets/'dopamine-wordmark.png', 7.03, 3.18, 2.55, 0.60)
    text(s, 7.10, 3.99, 2.52, 0.95, 'Meet Ava.\nYour in-store\nshopping advisor.', 16.5, PURPLE, True)
    note(s, 'Registered team identity supplied by the team leader. Product engineering is led by Muhammad Fikri Haikal. No completed contribution is attributed to the other member; QA and demo support are proposed next responsibilities, not a claim of prior coding work.')

    s = original[2]
    setup(s, 'Brief about the idea', numbering[id(s)], assets)
    text(s, 0.40, 1.65, 6.1, 1.14, 'A useful shopping conversation.\nInside the store.', 25, DARK, True)
    text(s, 0.43, 3.02, 5.90, 1.64,
         'Describe a need and budget in Bahasa Indonesia.\nAva explains catalog-grounded options and helps navigate the same storefront.\nCart changes require review and fresh confirmation; checkout is a marked sandbox.', 14.5)
    pill(s, 0.43, 4.85, 3.25, '80 DEMO PRODUCTS  /  7 CATEGORIES')
    picture(s, assets/'ava.jpg', 6.92, 1.52, 2.63, 3.57)
    note(s, 'Actual deployed avatar capture from the October 18 Gemini 3.8 release. The product is a demo storefront, not an active merchant or a promise of inventory/fulfillment. Voice and typed interaction exist; full spoken commerce is still being validated.')

    s = original[3]
    setup(s, 'Opportunities', numbering[id(s)], assets, 'A focused commerce workflow, rather than a detached conversation.')
    cards = [
        ('01  DIFFERENTIATION', 'Advice that acts', 'Catalog advice connects to visible filters, comparison and reviewed cart changes in the same shopping session.'),
        ('02  PROBLEM SOLVED', 'Clearer choices', 'Turn needs and a budget into a relevant shortlist, with product trade-offs and a concrete next step.'),
        ('03  USP', 'Verified next steps', 'Known catalog facts, budget-aware choices and action receipts keep answers tied to what the store actually did.'),
    ]
    for i, (label, heading, body) in enumerate(cards):
        x = 0.38 + i*3.13
        rect(s, x, 1.65, 3.0, 2.82, CREAM if i != 1 else LILAC, rounded=True)
        text(s, x+0.18, 1.88, 2.66, 0.27, label, 10, PURPLE, True)
        text(s, x+0.18, 2.42, 2.63, 0.7, heading, 21, DARK, True)
        text(s, x+0.18, 3.20, 2.62, 1.12, body, 13)
    text(s, 0.43, 4.70, 9.05, 0.42, 'Potential impact to validate: less discovery effort, better task completion and greater confidence. No sales or conversion uplift has been measured.', 11.5, MUTED)

    s = original[4]
    setup(s, 'List of features offered by the solution', numbering[id(s)], assets)
    features = [
        ('Catalog-aware advice', 'Uses approved page and product snapshots, not arbitrary screen pixels.'),
        ('Budget-based discovery', 'Shortlists demo products and explains recorded benefits and trade-offs.'),
        ('Visible store actions', 'Filter, navigate, highlight and compare 2–3 products from supported groups.'),
        ('Reviewed cart changes', 'Preview, fresh confirmation, browser receipt and guarded undo.'),
        ('Compact avatar controls', 'Microphone, optional Chat and Disconnect; browsing keeps the canvas.'),
        ('QRIS sandbox checkout', 'Server-priced preview, confirmed QR creation and verified simulator status.'),
    ]
    for i, (heading, body) in enumerate(features):
        x = 0.4 + (i % 3)*3.12
        y = 1.64 + (i // 3)*1.68
        rect(s, x, y, 2.98, 1.49, LILAC if i % 2 == 0 else CREAM, rounded=True)
        text(s, x+0.16, y+0.19, 2.67, 0.43, heading, 14.5, PURPLE, True)
        text(s, x+0.16, y+0.76, 2.66, 0.65, body, 11.8)
    text(s, 0.43, 5.05, 9.0, 0.19, 'Implemented prototype features. Cart, stock, prices and fulfillment are demo data; payments use Midtrans Sandbox.', 9, MUTED)

    s = original[5]
    setup(s, 'Process flow diagram or Use-case diagram', numbering[id(s)], assets, 'Implemented journey; end-to-end human voice validation remains pending.')
    stages = [
        ('01', 'Open Ava', 'Visible face and microphone readiness before room allocation.'),
        ('02', 'Share needs', 'Voice or optional typed input describes needs and budget.'),
        ('03', 'Explore', 'Filter, compare and review a relevant shortlist.'),
        ('04', 'Review cart', 'Inspect the proposed item and quantity change.'),
        ('05', 'Confirm', 'Fresh consent, applied action and a browser receipt.'),
        ('06', 'Sandbox', 'Review total; confirm QR; settle in the official simulator.'),
    ]
    for i, (n, heading, body) in enumerate(stages):
        x = 0.39 + i*1.56
        rect(s, x, 2.02, 1.42, 2.34, LILAC if i % 2 == 0 else CREAM, rounded=True)
        text(s, x+0.11, 2.2, 1.18, 0.35, n, 21, PURPLE, True)
        text(s, x+0.11, 2.85, 1.18, 0.46, heading, 13.5, DARK, True)
        text(s, x+0.11, 3.44, 1.18, 0.82, body, 10.4)
        if i < 5:
            arrow(s, x+1.43, 3.02, x+1.54, 3.02)
    rect(s, 0.42, 4.65, 9.10, 0.40, LIME, rounded=True)
    text(s, 0.56, 4.75, 8.83, 0.20, 'Mutation boundary: review → explicit confirmation → verified outcome. A timeout is not a success receipt.', 10.2, DARK, True)

    s = original[7]
    setup(s, 'Architecture diagram of the proposed solution', numbering[id(s)], assets)
    # Connectors go behind labeled native PowerPoint nodes.
    arrow(s, 2.50, 2.79, 3.04, 2.79, both=True)
    arrow(s, 6.63, 2.58, 7.03, 2.12, both=True)
    arrow(s, 6.63, 3.30, 7.03, 3.25, both=True)
    arrow(s, 6.63, 3.79, 7.03, 4.32, both=True)
    arrow(s, 4.84, 4.49, 4.84, 4.14)
    rect(s, 0.40, 2.06, 2.10, 1.44, CREAM, LINE, True)
    text(s, 0.56, 2.25, 1.79, 0.46, "Shopper's browser", 14, PURPLE, True)
    text(s, 0.56, 2.86, 1.77, 0.51, 'Store + React widget\nPage state / action receipts', 11)
    rect(s, 3.08, 1.65, 3.52, 2.51, LILAC, LINE, True)
    text(s, 3.26, 1.86, 3.16, 0.29, 'GOOGLE CLOUD RUN', 12.5, PURPLE, True)
    text(s, 3.26, 2.22, 3.10, 0.19, 'asia-southeast1 • one combined service', 9.2, MUTED)
    rect(s, 3.27, 2.66, 3.14, 0.51, WHITE, rounded=True)
    text(s, 3.4, 2.81, 2.85, 0.19, 'Node storefront + session / checkout APIs', 9.8, DARK, True)
    rect(s, 3.27, 3.42, 3.14, 0.49, WHITE, rounded=True)
    text(s, 3.4, 3.57, 2.84, 0.19, 'Python voice agent + bounded tools', 10.6, DARK, True)
    for y, heading, body in [
        (1.67, 'Gemini Live API', '3.8 native audio + tool selection'),
        (2.81, 'LiveKit + Spatius', 'RTC transport + avatar rendering'),
        (3.95, 'Midtrans Sandbox', 'Confirmed QRIS / verified status'),
    ]:
        rect(s, 7.07, y, 2.50, 0.82, CREAM, LINE, True)
        text(s, 7.22, y+0.14, 2.20, 0.26, heading, 13, PURPLE, True)
        text(s, 7.22, y+0.50, 2.20, 0.22, body, 9.2)
    rect(s, 3.31, 4.48, 3.08, 0.49, LIME, rounded=True)
    text(s, 3.45, 4.63, 2.80, 0.18, 'Secret Manager • runtime credentials', 10.2, DARK, True)
    text(s, 0.42, 4.18, 2.07, 0.72, 'Browser ↔ LiveKit carries RTC; the worker joins the same room. Spatius drives the visible avatar.', 9.5, MUTED)
    note(s, 'Functional service boundaries, not separate Cloud Run deployments. Browser code executes in the client; Cloud Run hosts its assets. Gemini uses the direct Google API. LiveKit/Spatius are external real-time providers. Sessions are in memory and service maximum remains one instance. Cart lives in the browser. Shared state is future scaling work.')

    s = original[8]
    setup(s, 'Technologies to be used in the solution', numbering[id(s)], assets, 'Deployed prototype stack; Google AI and Google Cloud are core runtime components.')
    tech = [
        ('GOOGLE AI + CLOUD', [('Gemini 3.8 Live', 'Native audio and 21 bounded tools with explicit blocking behavior.'), ('Cloud Run', 'Hosts the storefront, Node APIs and continuous Python worker.'), ('Secret Manager', 'Server credentials mounted only at runtime.')]),
        ('REALTIME EXPERIENCE', [('React + TypeScript', 'Compact Ava widget embedded in the Node multipage storefront.'), ('LiveKit', 'Rooms, audio transport and agent communication.'), ('Spatius AvatarKit', 'Visible rendering before voice-session allocation.')]),
        ('COMMERCE + DELIVERY', [('Midtrans Sandbox', 'QRIS testing with the official simulator; no real funds or orders.'), ('Node 22 + Python 3.12', 'Pinned dependencies and repeatable container builds.'), ('Signed receipts', 'Server-priced checkout and stable order IDs for retries.')]),
    ]
    for i, (label, rows) in enumerate(tech):
        x = 0.4 + i*3.12
        rect(s, x, 1.65, 2.98, 3.36, CREAM if i != 1 else LILAC, rounded=True)
        text(s, x+0.16, 1.84, 2.67, 0.23, label, 10.2, PURPLE, True)
        for j, (heading, body) in enumerate(rows):
            y = 2.27 + j*0.87
            text(s, x+0.16, y, 2.66, 0.24, heading, 12.5, DARK, True)
            text(s, x+0.16, y+0.34, 2.65, 0.42, body, 10.4)
    note(s, 'Firebase, Firestore and Vertex AI are not used and are not claimed. AvatarKit Web 1.3.10, RTC 1.1.5, Python Agents/Google/Spatius/Tavus plugins 1.8.8, browser LiveKit client 2.23.0. Tavus plugin is retained dependency support, not the active avatar provider. CI has not run on GitHub. Runtime worker/session limits are unchanged by the upgrade.')

    s = original[10]
    setup(s, 'Snapshots of the prototype', numbering[id(s)], assets, 'Actual deployed UI • Gemini 3.8 release • typed assistant input • English captions')
    picture(s, assets/'filter-checkout.jpg', 0.4, 1.68, 7.2, 3.42)
    for y, n, heading, body in [
        (1.76, '01', 'Filtered catalog', 'Headphones within IDR 500,000.'),
        (2.88, '02', 'Server-priced review', 'Basket and example shipping.'),
        (4.00, '03', 'Ava stays present', 'Chat closed. Same canvas and demo cart.'),
    ]:
        text(s, 7.84, y, 1.65, 0.17, n, 9, PURPLE, True)
        text(s, 7.84, y+0.23, 1.65, 0.40, heading, 12.2, DARK, True)
        text(s, 7.84, y+0.69, 1.65, 0.35, body, 9.3, MUTED)
    note(s, 'Actual October 18 deployed typed-tools QA. Both filter and checkout preview passed. Worker filter receipt applied in 472 ms; checkout review_ready in 585 ms. No QR/payment was initiated on this image. Screenshot is cropped to exclude the demo order identifier; no recipient data, credentials or join tokens are shown.')

    s = screenshot_extra
    setup(s, 'Snapshots of the prototype — sandbox checkout', numbering[id(s)], assets, 'Actual earlier provider test • official simulator • separate from the current-release preview')
    picture(s, assets/'sandbox-paid.jpg', 0.47, 1.62, 2.82, 3.53)
    text(s, 3.76, 1.82, 5.65, 0.72, 'A verified simulation,\nwith a visible result.', 22, DARK, True)
    steps = [('1', 'Review the server-priced basket'), ('2', 'Confirm and generate a sandbox QR'), ('3', 'Settle through the official simulator'), ('4', 'Read verified payment status in the UI')]
    for i, (n, label) in enumerate(steps):
        y = 2.97+i*0.41
        pill(s, 3.82, y, 0.32, n)
        text(s, 4.33, y+0.03, 5.0, 0.28, label, 12)
    text(s, 3.82, 4.82, 5.53, 0.32, 'Sandbox only. No real funds, order, shipping or inventory reservation.', 10.5, MUTED)
    note(s, 'Actual Midtrans QR/simulator settlement verification from the earlier dopamine-midtrans-20261018 deployment, October 18. Not rerun on the Gemini 3.8 image. Full spoken payment/success announcement is not verified. The screenshot excludes the demonstration order identifier and contains no recipient information.')

    s = original[11]
    setup(s, 'Prototype Performance report/Benchmarking', numbering[id(s)], assets, 'Recorded 18 October 2026 • test results and observed timings, not production guarantees')
    for i, (n, label) in enumerate([('90', 'Catalog + Node checks'), ('38', 'Python unit checks'), ('18', 'Offline SDK checks'), ('51', 'Browser case-runs¹')]):
        x = 0.4+i*2.34
        rect(s, x, 1.65, 2.18, 0.95, LILAC if i%2 == 0 else CREAM, rounded=True)
        text(s, x+0.15, 1.77, 1.87, 0.39, n, 25, PURPLE, True)
        text(s, x+0.15, 2.30, 1.88, 0.17, label, 9.8)
    text(s, 0.44, 2.94, 5.20, 0.29, 'Actual cloud observations', 15.5, DARK, True)
    for y, label, value in [(3.47, 'Typed filter receipt', '472 ms / applied'), (3.94, 'Typed checkout review receipt', '585 ms / review_ready'), (4.41, 'First audio → playback RPC²', '881–3870 ms / n=4')]:
        text(s, 0.44, y, 3.10, 0.26, label, 11.5)
        text(s, 3.61, y, 2.20, 0.26, value, 11.2, PURPLE, True)
    rect(s, 6.17, 2.93, 3.35, 1.87, PINK, rounded=True)
    text(s, 6.36, 3.11, 2.97, 0.28, 'Still to validate', 14, DARK, True)
    text(s, 6.36, 3.57, 2.95, 1.09, 'Human voice checkout and reliability.\nSynthetic voice filter did not pass.\nCold asset downloads remain heavy.\nNo controlled model latency A/B.', 10.8)
    text(s, 0.44, 4.91, 9.0, 0.30, '¹ 17 scenarios × 3 widths; simulated assistant transport. Seven renderer cases also pass.  ² Mixed sessions, not speech-to-speaker E2E.', 8.5, MUTED)
    note(s, 'Node: 3 catalog + 87 module, actual Node 22.23.3/npm 12.2.0 also passed clean install/build. Python: 38 unit, 18 offline SDK cases. Browser shopping: 17 scenarios at 1440/390/320; checkout uses a synthetic provider, renderer uses real public assets with session interception. Actual cloud timing samples are distinct, not a statistically representative benchmark. Voice input was synthetic; no human E2E or semantic latency pairing is claimed. The earlier 20-second receipt timeout remains unresolved. No measured conversion/revenue/customer figures.')

    s = original[12]
    setup(s, 'Additional Details/Future Development (if any)', numbering[id(s)], assets)
    roadmap = [
        ('NEXT', 'Validate the live experience', 'Rehearse human Indonesian speech, interruption, cart confirmation and spoken sandbox success. Measure end-to-end delay and investigate intermittent receipts.'),
        ('PRODUCTION PATH', 'Connect a real merchant', 'Integrate authorized inventory, orders and fulfillment. Add shared ownership and rate-limit state before scaling beyond one Cloud Run instance.'),
        ('BUSINESS VALIDATION', 'Measure shopper outcomes', 'Pilot with merchants and measure time-to-shortlist, task completion, checkout drop-off and shopper confidence. Commercial impact is not yet measured.'),
    ]
    for i, (label, heading, body) in enumerate(roadmap):
        x = 0.4+i*3.12
        rect(s, x, 1.65, 2.98, 2.97, CREAM if i%2 == 0 else LILAC, rounded=True)
        text(s, x+0.17, 1.85, 2.64, 0.27, label, 9.6, PURPLE, True)
        text(s, x+0.17, 2.30, 2.64, 0.66, heading, 18, DARK, True)
        text(s, x+0.17, 3.15, 2.64, 1.25, body, 11.5)
    text(s, 0.44, 4.84, 9.0, 0.31, 'Provenance: DOPAMINE adds a commerce experience to the Archava avatar/session starter. Upstream attribution is retained.', 10, MUTED)

    s = original[13]
    setup(s, 'Provide links to your:', numbering[id(s)], assets, 'Prototype and demo folder linked. Public repository access remains pending.')
    rows = [
        (1.70, 'FINAL PRODUCT', 'Open the deployed DOPAMINE prototype', URL, 'LIVE', LIME),
        (2.87, 'GITHUB PUBLIC REPOSITORY', 'Archzxy / archava-commerce', REPO, 'PRIVATE — ACCESS PENDING', PINK),
        (4.04, 'DEMO VIDEO LINK', 'Demo video — Google Drive folder', VIDEO, 'SHARED FOLDER', LIME),
    ]
    for y, label, heading, link, status, fill in rows:
        rect(s, 0.4, y, 9.12, 1.02, CREAM, rounded=True)
        text(s, 0.57, y+0.13, 5.9, 0.18, label, 9.4, PURPLE, True)
        text(s, 0.57, y+0.43, 5.9, 0.28, heading, 13.2, DARK, True, link=link)
        if link:
            text(s, 0.57, y+0.79, 7.50, 0.16, link, 8.8, MUTED, link=link)
        pill(s, 6.91, y+0.20, 2.39, status, fill)
    note(s, 'The team supplied the exact demo folder URL on 19 October 2026. An anonymous HTTP request returned 200, title Demo Video - Google Drive, with video MIME metadata and no sign-in redirect; actual playback and duration were not verified. The authenticated GitHub CLI still reports PRIVATE; public repository access remains pending. Generating this deck does not change visibility, push source, publish or submit to the competition. Portal PDF limit is 5 MB. Video duration conflicts between portal/template/public FAQ remain documented separately; no organizer clarification is claimed.')

    # Keep the final organizer Thank You slide unchanged.
    props = prs.core_properties
    props.title = 'DOPAMINE — Google Cloud AI Builder Cup 2026'
    props.subject = 'Retail & Commerce prototype pitch; demo folder linked, public repository access pending'
    props.author = 'Archzxy'
    props.keywords = 'DOPAMINE, Ava, Retail & Commerce, Gemini, Cloud Run, sandbox'
    props.comments = 'Prepared from the prescribed organizer template. Evidence checkpoint: 18 October 2026; demo folder link updated: 19 October 2026.'
    args.output.mkdir(parents=True, exist_ok=True)
    path = args.output / 'DOPAMINE_AI_Builder_Cup_2026.pptx'
    prs.save(path)
    assert len(prs.slides) == 13
    print(path)
    print('13 slides; organizer banner/header/footer/closing retained; all mandatory sections populated')


if __name__ == '__main__':
    main()
